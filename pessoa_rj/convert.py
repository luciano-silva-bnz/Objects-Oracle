# -*- coding: utf-8 -*-
"""
Extrai comprovantes PDF a partir de um relatório TXT e um PDF macro de comprovantes.
Gera arquivos individuais, Excel e ZIP na pasta escolhida.

Requisitos:
    pip install pypdf openpyxl
    # ou pip install PyPDF2 openpyxl
"""

import os
import re
import sys
import unicodedata
import zipfile
import warnings
from pathlib import Path
from typing import Iterable
from pessoa_rj.logging_config import get_logger

logger = get_logger(__name__)
# Preferência: pypdf > PyPDF2
try:
    from pypdf import PdfReader, PdfWriter
except ImportError:
    from PyPDF2 import PdfReader, PdfWriter

# Excel
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

import tkinter as tk
from tkinter import ttk, filedialog, messagebox

# Silencia warning do cryptography (opcional)
try:
    from cryptography.utils import CryptographyDeprecationWarning
    warnings.filterwarnings("ignore", category=CryptographyDeprecationWarning)
except Exception:
    pass

STOPWORDS = {'DE', 'DA', 'DO', 'DAS', 'DOS', 'E', 'A', 'O', 'U'}

def norm_text(s: str) -> str:
    """Normaliza e limpa o texto (remove acentos, maiúsculo, só letras/números/espaco)."""
    s = unicodedata.normalize("NFKD", s or "")
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = s.upper()
    s = re.sub(r"[^A-Z0-9 ]+", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s

def read_pdf_texts(pdf_path: str):
    """Extrai o texto de cada página do PDF."""
    reader = PdfReader(pdf_path)
    texts = []
    for page in reader.pages:
        try:
            texts.append(page.extract_text() or "")
        except Exception:
            texts.append("")
    return texts

def safe_filename(name: str) -> str:
    """
    Mantém apenas [A-Za-z0-9-_.() ] e troca / e \ por '-'.
    NÃO troca espaço por underscore; quem chama decide o formato.
    """
    base = name or "sem_nome"
    base = base.replace("/", "-").replace("\\", "-")
    base = re.sub(r"[^A-Za-z0-9\-_.() ]+", "", base)
    base = re.sub(r"\s{2,}", " ", base).strip()
    # limpa hifens/underscores duplicados ocasionais
    base = re.sub(r"-{2,}", "-", base)
    base = re.sub(r"_{2,}", "_", base)
    return base[:180] if base else "sem_nome"

def fold_upper_keep_punct(s: str) -> str:
    """
    Remove acentos e sobe pra MAIÚSCULAS, mas **mantém pontuação** (ex.: '/')
    e comprime espaços. Útil pra procurar 'DATA DO CREDITO' e 'DD/MM/AAAA'.
    """
    s = unicodedata.normalize("NFKD", s or "")
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = s.upper()
    s = re.sub(r"\s+", " ", s)
    return s

def group_contiguous(pages):
    """Agrupa páginas contíguas em ranges."""
    if not pages: return []
    pages = sorted(set(pages))
    groups = []; start = prev = pages[0]
    for p in pages[1:]:
        if p == prev + 1:
            prev = p
        else:
            groups.append((start, prev))
            start = prev = p
    groups.append((start, prev))
    return groups

def ler_pessoas_do_relatorio_txt(txt_path: str):
    """
    Lê o TXT separado por ';' e extrai (codigo, nome_completo) do primeiro campo ("Pessoa").
    Aceita linhas do tipo: "30491 - DANIELE DE MELO VIEIRA;..."
    """
    registros = []
    with open(txt_path, "r", encoding="utf-8", errors="ignore") as f:
        for line in f:
            row = line.rstrip("\n")
            if not row.strip():
                continue
            if row.upper().startswith("PESSOA;"):
                continue  # cabeçalho
            partes = row.split(";")
            if not partes:
                continue
            campo_pessoa = partes[0].strip()
            codigo = ""
            nome = campo_pessoa
            if "-" in campo_pessoa:
                left, right = campo_pessoa.split("-", 1)
                codigo = re.sub(r"\D", "", left).strip()
                nome = re.sub(r"\s+", " ", right.strip())
            else:
                # fallback: tenta pegar o primeiro número como código
                m = re.match(r"\s*(\d+)\s+(.*)", campo_pessoa)
                if m:
                    codigo = m.group(1)
                    nome = re.sub(r"\s+", " ", (m.group(2) or "").strip())
            if nome:
                registros.append({"codigo": codigo, "nome": nome})
    return registros

def compila_padrao_nome_completo(nome_norm: str):
    """
    Compila regex para o NOME COMPLETO normalizado como palavras inteiras,
    tolerando múltiplos espaços entre palavras.
    """
    tokens = nome_norm.split()
    if not tokens:
        return None
    inner = r"\s+".join(map(re.escape, tokens))
    return re.compile(rf"\b{inner}\b")

# --- Extração da "Data do Crédito" no texto do comprovante -------------------

_PATS_CREDITO = [
    r"DATA\s*DO\s*CR[ÉE]DITO[:\s]*([0-3]\d/[01]\d/\d{4})",
    r"DT\.?\s*CR[ÉE]DITO[:\s]*([0-3]\d/[01]\d/\d{4})",
    r"DATA\s*CR[ÉE]DITO[:\s]*([0-3]\d/[01]\d/\d{4})",
    r"CR[ÉE]DITO[:\s]*([0-3]\d/[01]\d/\d{4})",
]

# --- extração da Data do Crédito --------------------------------------------

_DATE_ANY = r"([0-3]\d/[01]\d/\d{4})"

def extrair_data_credito(texto_pagina: str) -> str | None:
    """
    1) Procura 'DATA DO CREDITO' próximo da data (mesma linha/coluna),
    2) tenta uma janela em torno de 'CREDITO',
    3) ou cai na primeira data da página.
    Tudo com texto MAIÚSCULO sem acento, **preservando '/'**.
    """
    T = fold_upper_keep_punct(texto_pagina)

    # 1) rótulo explícito
    m = re.search(r"DATA\s*DO\s*CREDITO[^\n\r]{0,40}?" + _DATE_ANY, T)
    if m:
        return m.group(1)

    # 2) janela em torno da palavra CREDITO
    for mm in re.finditer(r"CREDITO", T):
        start = max(0, mm.start() - 100)
        end = min(len(T), mm.end() + 100)
        win = T[start:end]
        m2 = re.search(_DATE_ANY, win)
        if m2:
            return m2.group(1)

    # 3) fallback: primeira data da página
    m3 = re.search(_DATE_ANY, T)
    return m3.group(1) if m3 else None
# --- Excel --------------------------------------------------------------------

def escrever_excel(resumo_rows, xlsx_path: Path):
    """
    Escreve planilha Excel com colunas:
    [codigo, nome_do_relatorio, paginas_macro, data_credito, arquivo_gerado, encontrado, arquivo_macro]
    """
    wb = Workbook()
    ws: Worksheet = wb.active
    ws.title = "resumo"

    headers = ["codigo", "nome_do_relatorio", "paginas_macro", "data_credito",
               "arquivo_gerado", "encontrado", "arquivo_macro"]
    ws.append(headers)

    for row in resumo_rows:
        ws.append(list(row))

    widths = [12, 42, 16, 16, 80, 12, 28]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    wb.save(xlsx_path)


# --- Núcleo -------------------------------------------------------------------

def extrair_comprovantes(relatorio_txt: str, macro_pdf_paths: str | Iterable[str], pasta_saida: str) -> dict:
    """
    Aceita 1 ou N PDFs macro (str ou lista de str). Todos serão varridos.
    """
    # normaliza para lista
    if isinstance(macro_pdf_paths, (str, Path)):
        macro_list = [str(macro_pdf_paths)]
    else:
        macro_list = [str(p) for p in macro_pdf_paths]

    out_dir = Path(pasta_saida)
    out_dir.mkdir(parents=True, exist_ok=True)
    logger.info('Iniciando extração: txt=%s, pdfs=%d, saída=%s', relatorio_txt, len(macro_list), out_dir)

    # 1) Ler (codigo, nome) do TXT e pré-calcular expressões
    pessoas = ler_pessoas_do_relatorio_txt(relatorio_txt)
    nomes_norm = [norm_text(p["nome"]) for p in pessoas]
    patterns = [compila_padrao_nome_completo(n) if n else None for n in nomes_norm]

    resumo_rows = []
    total_gerados = 0

    # Itera por cada PDF macro informado
    for macro_pdf in macro_list:
        macro_reader = PdfReader(macro_pdf)
        macro_base = Path(macro_pdf).stem
        macro_textos = read_pdf_texts(macro_pdf)
        macro_norm = [norm_text(t) for t in macro_textos]

        for i, pessoa in enumerate(pessoas):
            codigo = pessoa.get("codigo", "")
            nome_original = pessoa["nome"]
            nome_norm = nomes_norm[i]
            pad = patterns[i]

            hits = []
            if pad is not None and nome_norm:
                for idx_pag, page_norm in enumerate(macro_norm):
                    if pad.search(page_norm):
                        hits.append(idx_pag)

            if hits:
                for (a, b) in group_contiguous(hits):
                    data_credito = extrair_data_credito(macro_textos[a]) or ""
                    data_credito_fn = data_credito.replace("/", "-") if data_credito else "DATA-NAO-ENCONTRADA"
                    paginas_str = f"{a+1}" if a == b else f"{a+1}-{b+1}"
                    nome_slug = re.sub(r"\s+", "_", nome_original.strip())

                    # inclui o nome do arquivo macro na linha de resumo para rastreabilidade

                    fname_humano = f"{codigo}-{nome_slug}-{paginas_str}-({data_credito_fn})"
                    fname = safe_filename(fname_humano) + ".pdf"
                    fpath = out_dir / fname

                    writer = PdfWriter()
                    for p in range(a, b + 1):
                        writer.add_page(macro_reader.pages[p])
                    with open(fpath, "wb") as f:
                        writer.write(f)

                    resumo_rows.append((
                        codigo,
                        nome_original,
                        paginas_str,
                        data_credito,
                        str(fpath),
                        True,
                        macro_base   # <<< nome do PDF de origem
                    ))
                    total_gerados += 1
            else:
                # registra ausência *para esse macro* apenas uma vez por pessoa
                resumo_rows.append((
                    codigo,
                    nome_original,
                    "",
                    "",
                    "",
                    False,
                    macro_base   # <<< nome do PDF de origem
                ))


    # 4) Excel + ZIP (ZIP SÓ COM PDFs; salvo na MESMA PASTA de saída)
    xlsx_path = out_dir / "resumo_extracao.xlsx"
    escrever_excel(resumo_rows, xlsx_path)

    zip_path = out_dir / "comprovantes.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for row in resumo_rows:
            path = row[4]  # arquivo_gerado
            ok = row[5]    # encontrado (True/False)
            if ok and path and Path(path).suffix.lower() == ".pdf":
                zf.write(path, arcname=Path(path).name)


    summary = {
        "total_registros_no_relatorio": len(pessoas),
        "total_comprovantes_gerados": total_gerados,
        "pasta_saida": str(out_dir.resolve()),
        "excel": str(xlsx_path.resolve()),
        "zip": str(zip_path.resolve()),
    }
    logger.info('Extração concluída: %s', summary)
    return summary
# --- GUI ----------------------------------------------------------------------

class App(tk.Tk):
    """Janela principal do extrator de comprovantes PDF (relatório TXT)."""
    def __init__(self):
        super().__init__()
        self.title("Extrair Comprovantes")
        self.geometry("780x320")
        self.minsize(640, 240)

        self.relatorio_path = tk.StringVar()
        self.macro_path = tk.StringVar()
        self.saida_dir = tk.StringVar(value=str(Path.cwd() / "comprovantes_out"))
        self.status = tk.StringVar(value="Selecione os arquivos e clique em GERAR.")
        self._macro_list = []  # lista real de PDFs macro selecionados

        pad = {"padx": 12, "pady": 8}

        ttk.Button(self, text="Relatório (TXT)...", command=self.pick_relatorio).grid(row=0, column=0, sticky="w", **pad)
        ttk.Entry(self, textvariable=self.relatorio_path).grid(row=0, column=1, columnspan=3, sticky="we", **pad)
        ttk.Button(self, text="Comprovantes (PDF macro)...", command=self.pick_macro).grid(row=1, column=0, sticky="w", **pad)
        ttk.Entry(self, textvariable=self.macro_path).grid(row=1, column=1, columnspan=3, sticky="we", **pad)
        ttk.Button(self, text="Pasta de saída...", command=self.pick_saida).grid(row=2, column=0, sticky="w", **pad)
        ttk.Entry(self, textvariable=self.saida_dir).grid(row=2, column=1, columnspan=3, sticky="we", **pad)
        ttk.Label(self, textvariable=self.status, foreground="#444").grid(row=3, column=0, columnspan=4, sticky="w", **pad)
        ttk.Button(self, text="GERAR", command=self.on_gerar).grid(row=4, column=2, sticky="e", **pad)
        ttk.Button(self, text="FECHAR", command=self.destroy).grid(row=4, column=3, sticky="w", **pad)

        for c in range(1, 4):
            self.columnconfigure(c, weight=1)
        self.after(100, self.centralizar)

    def centralizar(self):
        self.update_idletasks()
        w = self.winfo_width()
        h = self.winfo_height()
        sw = self.winfo_screenwidth()
        sh = self.winfo_screenheight()
        x = (sw - w) // 2
        y = (sh - h) // 2
        self.geometry(f"{w}x{h}+{x}+{y}")

    def pick_relatorio(self):
        p = filedialog.askopenfilename(title="Selecione o RELATÓRIO (TXT)", filetypes=[("TXT", "*.txt")])
        if p: self.relatorio_path.set(p)

    def pick_macro(self):
        paths = filedialog.askopenfilenames(
            title="Selecione um ou mais PDFs MACRO de COMPROVANTES",
            filetypes=[("PDF", "*.pdf")]
        )
        if paths:
            # guarda como string separada por ; só pra exibir no Entry
            self.macro_path.set("; ".join(paths))
            # e mantém também a lista real para o processamento
            self._macro_list = list(paths)


    def pick_saida(self):
        d = filedialog.askdirectory(title="Selecione a pasta de saída")
        if d: self.saida_dir.set(d)

    def on_gerar(self):
        rel = self.relatorio_path.get().strip()
        out = self.saida_dir.get().strip() or str(Path.cwd() / "comprovantes_out")

        # usa a lista real se houver; senão tenta ler um único caminho do Entry
        macro_list = getattr(self, "_macro_list", []) or [
            p.strip() for p in self.macro_path.get().split(";") if p.strip()
        ]

        if not rel or not macro_list:
            logger.warning('Seleção inválida: relatorio=%s, pdfs=%d', rel, len(macro_list))
            messagebox.showwarning("Faltando arquivo", "Selecione o RELATÓRIO (TXT) e pelo menos um PDF de COMPROVANTES.")
            return
        if not os.path.isfile(rel) or not rel.lower().endswith(".txt"):
            logger.warning('Relatório inválido selecionado: %s', rel)
            messagebox.showerror("Arquivo inválido", "O RELATÓRIO deve ser um arquivo .TXT válido.")
            return
        bad = [p for p in macro_list if not (os.path.isfile(p) and p.lower().endswith(".pdf"))]
        if bad:
            logger.warning('PDFs inválidos informados: %s', bad)
            messagebox.showerror("Arquivo inválido", f"Estes caminhos não são PDFs válidos:\n\n" + "\n".join(bad))
            return

        try:
            self.status.set(f"Processando {len(macro_list)} PDF(s)... aguarde.")
            self.update_idletasks()
            result = extrair_comprovantes(rel, macro_list, out)   # <— agora passa lista
            self.status.set("Concluído.")
            logger.info('Processamento finalizado para %s: %s', rel, result)
            msg = (
                "Concluído!\n\n"
                f"Registros no relatório: {result['total_registros_no_relatorio']}\n"
                f"Comprovantes gerados: {result['total_comprovantes_gerados']}\n\n"
                f"Pasta: {result['pasta_saida']}\nExcel: {result['excel']}\nZIP: {result['zip']}"
            )
            if messagebox.askyesno("Sucesso", msg + "\n\nAbrir a pasta de saída?"):
                try:
                    if sys.platform.startswith("win"):
                        os.startfile(result["pasta_saida"])
                    elif sys.platform == "darwin":
                        os.system(f'open "{result["pasta_saida"]}"')
                    else:
                        os.system(f'xdg-open "{result["pasta_saida"]}"')
                except Exception:
                    pass
        except Exception as e:
            self.status.set("Erro.")
            logger.exception('Erro ao gerar comprovantes')
            messagebox.showerror("Erro ao gerar", f"{type(e).__name__}: {e}")


def main():
    # Sua App já herda de tk.Tk e constrói toda a UI no __init__
    logger.info('Iniciando extrator de comprovantes (GUI)')
    app = App()
    # (opcional) um tema/estilo básico
    try:
        style = ttk.Style()
        style.theme_use('vista' if 'vista' in style.theme_names() else 'clam')
        style.configure('TButton', padding=(6, 2))
    except Exception:
        pass
    app.mainloop()
    logger.info('Extrator de comprovantes encerrado')
    return 0

if __name__ == '__main__':
    sys.exit(main())


