import pandas as pd
import tkinter as tk
from tkinter import filedialog
import tkinter.messagebox
import openpyxl

# Função para selecionar arquivos
def selecionar_arquivos():
    root = tk.Tk()
    root.withdraw()  # Esconde a janela principal do Tkinter
    file_paths = filedialog.askopenfilenames(title='Selecione as três planilhas do Excel')
    return list(file_paths)

# Função para selecionar o local e o nome do arquivo para salvar a nova planilha
def salvar_como():
    root = tk.Tk()
    root.withdraw()  # Esconde a janela principal do Tkinter
    file_path = filedialog.asksaveasfilename(defaultextension=".xlsx",
                                             filetypes=[("Excel files", "*.xlsx")],
                                             title="Salvar a nova planilha como",
                                             initialfile="analise de movimentação.xlsx")
    return file_path

# Seleciona os arquivos usando a interface gráfica
file_paths = selecionar_arquivos()

# Lista para armazenar os dataframes filtrados
dfs = []

# Lê cada arquivo do Excel e adiciona o dataframe filtrado à lista
for file in file_paths:
    # Lê o arquivo do Excel
    df = pd.read_excel(file, sheet_name='Exportar Planilha', dtype=str)
    # Verifica se as colunas 'B9_FILIAL' e 'B9_LOCAL' existem antes de filtrar
    if 'B9_FILIAL' in df.columns and 'B9_LOCAL' in df.columns:
        # Filtra o dataframe pelas condições especificadas
        df_filtered = df[(df['B9_FILIAL'] == '41') & (df['B9_LOCAL'] == '01')]
        # Adiciona o dataframe filtrado à lista se não estiver vazio
        if not df_filtered.empty:
            dfs.append(df_filtered)
    else:
        # Se as colunas 'B9_FILIAL' e 'B9_LOCAL' não existirem, adiciona o dataframe inteiro à lista
        dfs.append(df)

if dfs:
    final_df = pd.concat(dfs, ignore_index=True)
else:
    print("Nenhum dataframe para concatenar. Verifique os arquivos e condições de filtragem.")
    # Pode adicionar mais lógica aqui, como uma saída antecipada do script ou outra ação

# Cria um novo dataframe para armazenar os valores de código e descrição do produto
new_df = pd.DataFrame(columns=['Codigo', 'Descrição'])

# Procura os campos especificados em cada planilha e adiciona os valores ao novo dataframe
for codigo_col in ['COD PRODUTO', 'B9_COD']:
    if codigo_col in final_df.columns:
        for descricao_col in ['DESCRICAO', 'B1_DESC']:
            if descricao_col in final_df.columns:
                # Extrai os dados das colunas código e descrição
                temp_df = final_df[[codigo_col, descricao_col]].drop_duplicates()
                temp_df.columns = ['Codigo', 'Descrição']  # Renomeia as colunas
                # Verifica se o código e a descrição estão presentes antes de adicionar ao novo dataframe
                for index, row in temp_df.iterrows():
                    if pd.notnull(row['Codigo']) and pd.notnull(row['Descrição']):
                        if row['Codigo'] not in new_df['Codigo'].values:
                            new_df = pd.concat([new_df, pd.DataFrame([row])], ignore_index=True)

####################################
#Atualizar a planilha de comparação#
####################################

# Função para selecionar um arquivo
def selecionar_arquivo():
    root = tk.Tk()
    root.withdraw()  # Esconde a janela principal do Tkinter
    file_path = filedialog.askopenfilename(title='Selecione a planilha do Excel')
    return file_path

# Seleciona o arquivo usando a interface gráfica
file_path = selecionar_arquivo()

# Carrega a planilha do Excel
wb = openpyxl.load_workbook(file_path)

# Verifica se a aba "Comparativo" existe
if 'Comparativo' in wb.sheetnames:
    ws = wb['Comparativo']
else:
    # Cria a aba "Comparativo" se ela não existir
    ws = wb.create_sheet('Comparativo')

# Limpa as colunas A e B a partir da linha 3
for i in range(3, ws.max_row + 1):
    ws[f'A{i}'] = None
    ws[f'B{i}'] = None

# Escreve os dados no arquivo do Excel a partir da célula A3
row_num = 3
for index, row in new_df.iterrows():
    ws.cell(row=row_num, column=1, value=row['Codigo'])
    ws.cell(row=row_num, column=2, value=row['Descrição'])
    row_num += 1

# Salva a planilha com o mesmo nome
wb.save(file_path)

# Exibe uma caixa de diálogo informando que a planilha foi atualizada
tkinter.messagebox.showinfo("Informação", f"A planilha foi atualizada com sucesso: {file_path}")

print(f"A planilha foi atualizada com sucesso: {file_path}")

